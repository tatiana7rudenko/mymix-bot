"""Минимальный конвертер Markdown -> DOCX для фактуры и пособия.

Поддерживает: # / ## / ### заголовки, абзацы, списки (- и 1.), > цитаты-врезки,
**жирный**, *курсив*, [ссылки](url), голые URL, --- как разрыв страницы, простые таблицы.
Использование: python3 md2docx.py input.md output.docx [--accent HEX]
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ACCENT = "8A3B5C"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\)|https?://[^\s)]+)")


def add_hyperlink(par, url, text):
    part = par.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FA8")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    par._p.append(link)


def add_inline(par, text, base_bold=False, base_italic=False):
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            add_inline(par, piece[2:-2], True, base_italic)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            add_inline(par, piece[1:-1], base_bold, True)
        elif piece.startswith("[") and re.match(r"\[([^\]]+)\]\(([^)]+)\)", piece):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", piece)
            add_hyperlink(par, m.group(2), m.group(1))
        elif piece.startswith("http"):
            add_hyperlink(par, piece, piece)
        else:
            r = par.add_run(piece)
            r.bold = base_bold or None
            r.italic = base_italic or None


def shade(par, fill):
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("w:val", "single"), ("w:sz", "24"), ("w:space", "8"), ("w:color", ACCENT)):
        left.set(qn(k), v)
    bdr.append(left)
    ppr.append(bdr)


def restart_numbering(doc, par):
    """Каждый нумерованный список начинается с 1: новый w:num с startOverride."""
    numbering = doc.part.numbering_part.element
    style_num = doc.styles["List Number"].element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(style_num).abstractNumId.val
    num = numbering.add_num(abstract)
    ov = num.add_lvlOverride(ilvl=0)
    ov.add_startOverride(1)
    numPr = par._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = 0
    numPr.get_or_add_numId().val = num.numId


def build(src, dst):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Georgia"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.2
    for name, size in (("Title", 26), ("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12.5)):
        s = doc.styles[name]
        s.font.name = "Arial"
        s.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(ACCENT if name != "Heading 3" else "222222")
        s.paragraph_format.space_before = Pt(14 if name != "Heading 3" else 10)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True

    lines = open(src, encoding="utf-8").read().splitlines()
    i = 0
    cur_num = None
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.strip() == "---":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif line.startswith("# "):
            doc.add_paragraph(line[2:].strip(), style="Title")
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            h = doc.add_heading("", level=3)
            add_inline(h, line[5:].strip())
        elif line.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].startswith(">"):
                buf.append(lines[i][1:].strip())
                i += 1
            for chunk in "\n".join(buf).split("\n\n"):
                p = doc.add_paragraph()
                shade(p, "F7EEF2")
                p.paragraph_format.left_indent = Cm(0.3)
                add_inline(p, " ".join(chunk.split("\n")))
            continue
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = t.cell(r, c)
                    cell.text = ""
                    add_inline(cell.paragraphs[0], val, base_bold=(r == 0))
            doc.add_paragraph()
            continue
        elif re.match(r"^\s*[-*] ", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+[.)] ", line):
            p = doc.add_paragraph(style="List Number")
            if cur_num is None or not re.match(r"^\s*\d+[.)] ", lines[i - 1] if i else ""):
                restart_numbering(doc, p)
                cur_num = p._p.pPr.numPr.numId.val
            else:
                numPr = p._p.get_or_add_pPr().get_or_add_numPr()
                numPr.get_or_add_ilvl().val = 0
                numPr.get_or_add_numId().val = cur_num
            add_inline(p, re.sub(r"^\s*\d+[.)] ", "", line))
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        i += 1
    doc.save(dst)


if __name__ == "__main__":
    if "--accent" in sys.argv:
        ACCENT = sys.argv[sys.argv.index("--accent") + 1]
    build(sys.argv[1], sys.argv[2])
